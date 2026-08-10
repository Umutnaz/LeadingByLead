import openpyxl
import json

wb = openpyxl.load_workbook('Følgeskab_parameter_model.xlsx')

def read_sheet_data(ws, start_row=2):
    headers = [cell.value for cell in ws[1]]
    data = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if all(v is None for v in row):
            break
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        data.append(row_dict)
    return data

# Read all relevant data
start_values = read_sheet_data(wb[wb.sheetnames[1]])  # Startvrdier
gears = read_sheet_data(wb[wb.sheetnames[3]])  # Gear
action_cards = read_sheet_data(wb[wb.sheetnames[4]])  # Action cards
phase_events = read_sheet_data(wb[wb.sheetnames[2]])  # Fasehndelser

# Create character ID mapping
char_id_mapping = {
    'D': 'daniel',
    'C': 'charlie',
    'I': 'inzo',
    'S': 'søren',
    'IS': 'ida',
    'L': 'lene'
}

char_name_mapping = {
    'D': ('Daniel', 'GV1'),
    'C': ('Charlie', 'LMG1'),
    'I': ('Inzo', 'GV3'),
    'S': ('Søren', 'LMG2'),
    'IS': ('Ida-Sofie', 'GV2'),
    'L': ('Lene', 'GV4')
}

# Generate C# code for starting values
print("=== STARTING VALUES C# CODE ===\n")

csharp_code = """
private static List<Character> CreateCharacters()
{
    return new List<Character>
    {
"""

for person in start_values:
    person_id = person['ID']
    name, role = char_name_mapping[person_id]
    tjenestemotivation = person['Tjenestemotivation']
    sociallyst = person['Sociale lyst']
    tillid = person['Tillid']
    stress = person['Stress']
    
    csharp_code += f"""        new()
        {{
            Name = "{name} - {role}",
            Description = "",
            BaseStats = new BaseStats
            {{
                TjenesteMotivation = {tjenestemotivation},
                Sociallyst = {sociallyst},
                Tillid = {tillid},
                Stress = {stress}
            }}
        }},
"""

csharp_code += """    };
}
"""

print(csharp_code)

# Organize gears by phase
print("\n=== GEARS DATA SUMMARY ===\n")
gears_by_phase = {}
for gear in gears:
    phase = gear['Fase']
    if phase not in gears_by_phase:
        gears_by_phase[phase] = []
    gears_by_phase[phase].append(gear)

for phase in ['Opstart', 'Drift', 'Overdragelse']:
    print(f"\n{phase}: {len(gears_by_phase.get(phase, []))} impacts")
    
# Organize action cards by phase
print("\n=== ACTION CARDS DATA SUMMARY ===\n")
cards_by_phase = {}
for card in action_cards:
    phase = card['Fase']
    if phase not in cards_by_phase:
        cards_by_phase[phase] = []
    cards_by_phase[phase].append(card)

for phase in ['Opstart', 'Drift', 'Overdragelse']:
    print(f"\n{phase}: {len(cards_by_phase.get(phase, []))} impacts")

# Save organized data to JSON for next step
with open('seeddata_organized.json', 'w', encoding='utf-8') as f:
    json.dump({
        'start_values': start_values,
        'gears_by_phase': {k: v for k, v in gears_by_phase.items()},
        'cards_by_phase': {k: v for k, v in cards_by_phase.items()},
        'phase_events': phase_events,
        'char_id_mapping': char_id_mapping,
        'char_name_mapping': {k: list(v) for k, v in char_name_mapping.items()}
    }, f, ensure_ascii=False, indent=2)

print("\n\nData organized and saved to seeddata_organized.json")

# Verification counts
print("\n=== VERIFICATION COUNTS ===")
print(f"Starting profiles: {len(start_values)}")
print(f"Gear impacts: {len(gears)}")
print(f"Action card impacts: {len(action_cards)}")
print(f"Phase events: {len(phase_events)}")


