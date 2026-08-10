import json

# Load organized data
with open('seeddata_organized.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

start_values = data['start_values']
phase_events = data['phase_events']
char_id_mapping = data['char_id_mapping']
char_name_mapping = data['char_name_mapping']

# Load Excel data again for gears and action cards
import openpyxl
wb = openpyxl.load_workbook('Følgeskab_parameter_model.xlsx')

def read_sheet_data(ws, start_row=2):
    headers = [cell.value for cell in ws[1]]
    data_list = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if all(v is None for v in row):
            break
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        data_list.append(row_dict)
    return data_list

gears = read_sheet_data(wb[wb.sheetnames[3]])  # Gear
action_cards = read_sheet_data(wb[wb.sheetnames[4]])  # Action cards

# Map person ID to character object references
def get_char_ref(person_id):
    name_role = char_name_mapping.get(person_id)
    if not name_role:
        return None
    name = name_role[0]
    # Find character by name
    name_lookup = {
        'Daniel': 'daniel',
        'Charlie': 'charlie',
        'Inzo': 'inzo',
        'Søren': 'søren',
        'Ida-Sofie': 'ida',
        'Lene': 'lene'
    }
    return name_lookup.get(name)

# Generate effect code line
def gen_effect_code(person_id, gear_or_card):
    char_ref = get_char_ref(person_id)
    if not char_ref:
        return None
    
    tjenestemotivation = gear_or_card.get('Δ Tjenestemotivation', 0) or 0
    sociallyst = gear_or_card.get('Δ Sociale lyst', 0) or 0
    tillid = gear_or_card.get('Δ Tillid', 0) or 0
    stress = gear_or_card.get('Δ Stress', 0) or 0
    reaction = gear_or_card.get('Reaktion', '')
    
    changes = []
    if tjenestemotivation != 0:
        changes.append(f'Change(nameof(CurrentStats.TjenesteMotivation), {tjenestemotivation})')
    if sociallyst != 0:
        changes.append(f'Change(nameof(CurrentStats.Sociallyst), {sociallyst})')
    if tillid != 0:
        changes.append(f'Change(nameof(CurrentStats.Tillid), {tillid})')
    if stress != 0:
        changes.append(f'Change(nameof(CurrentStats.Stress), {stress})')
    
    if not changes:
        return None
    
    reaction_escaped = reaction.replace('"', '\\"').replace('\n', '\\n') if reaction else ''
    effect_code = f'Effect({char_ref}, "{reaction_escaped}", {", ".join(changes)})'
    
    return effect_code

# Generate gear code for one option
def gen_gear_option_code(gears_for_option):
    effects = []
    for gear in gears_for_option:
        effect = gen_effect_code(gear['ID'], gear)
        if effect:
            effects.append(f'                    {effect}')
    
    if not effects:
        return None
    
    option_text = gears_for_option[0].get('Gearnavn', 'Option')
    description = gears_for_option[0].get('Beskrivelse', '')
    description_escaped = description.replace('"', '\\"').replace('\n', '\\n') if description else ''
    
    # Build option code with proper escaping
    newline_n = '\\n'
    effects_str = f',{newline_n}                    '.join(effects)
    option_code = f'''                Option(
                    "{option_text}{newline_n}{newline_n}{description_escaped}",

                    {effects_str}
                ),
'''
    return option_code

# Group gears by phase and number
gears_by_phase_and_num = {}
for gear in gears:
    phase = gear['Fase']
    num = gear['Gear']
    key = f"{phase}_{num}"
    if key not in gears_by_phase_and_num:
        gears_by_phase_and_num[key] = []
    gears_by_phase_and_num[key].append(gear)

print("=== GEAR OPTIONS CODE ===\n")
for phase in ['Opstart', 'Drift', 'Overdragelse']:
    print(f"\n// {phase} Phase:")
    for num in [1, 2, 3, 4]:
        key = f"{phase}_{num}"
        if key in gears_by_phase_and_num:
            gear_group = gears_by_phase_and_num[key]
            option_code = gen_gear_option_code(gear_group)
            if option_code:
                print(option_code)

# Group action cards by phase and number
cards_by_phase_and_num = {}
for card in action_cards:
    phase = card['Fase']
    num = card['Kort']
    key = f"{phase}_{num}"
    if key not in cards_by_phase_and_num:
        cards_by_phase_and_num[key] = []
    cards_by_phase_and_num[key].append(card)

print("\n\n=== ACTION CARDS CODE ===\n")
for phase in ['Opstart', 'Drift', 'Overdragelse']:
    print(f"\n// {phase} Phase:")
    for num in range(1, 9):
        key = f"{phase}_{num}"
        if key in cards_by_phase_and_num:
            card_group = cards_by_phase_and_num[key]
            option_code = gen_gear_option_code(card_group)
            if option_code:
                print(option_code)

print("\n\n=== PHASE EVENTS CODE ===\n")
for event in phase_events:
    phase = event.get('Fase', '')
    person_id = event.get('ID', '')
    tjenesteMotivation = event.get('Δ Tjenestemotivation', 0) or 0
    sociallyst = event.get('Δ Sociale lyst', 0) or 0
    tillid = event.get('Δ Tillid', 0) or 0
    stress = event.get('Δ Stress', 0) or 0
    begrundelse = event.get('Begrundelse', '')
    
    print(f"// {phase} - {person_id}: {tjenestemotivation},{sociallyst},{tillid},{stress}")


