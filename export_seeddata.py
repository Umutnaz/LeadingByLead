import openpyxl
import json
import os

wb = openpyxl.load_workbook('Følgeskab_parameter_model.xlsx')

def read_sheet_data(ws, start_row=2):
    headers = [cell.value for cell in ws[1]]
    data = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if all(v is None for v in row):
            break
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        data.append(row_dict)
    return data

# Read all data
start_values = read_sheet_data(wb[wb.sheetnames[1]])  # Startvrdier
gears = read_sheet_data(wb[wb.sheetnames[3]])  # Gear
action_cards = read_sheet_data(wb[wb.sheetnames[4]])  # Action cards
phase_events = read_sheet_data(wb[wb.sheetnames[2]])  # Fasehndelser

# Character mappings
char_names = {
    'D': 'Daniel',
    'C': 'Charlie',
    'I': 'Inzo',
    'S': 'Søren',
    'IS': 'Ida-Sofie',
    'L': 'Lene'
}

char_refs = {
    'D': 'daniel',
    'C': 'charlie',
    'I': 'inzo',
    'S': 'søren',
    'IS': 'ida',
    'L': 'lene'
}

# Save to JSON for other uses
data_export = {
    'characters': char_names,
    'gears': gears,
    'action_cards': action_cards,
    'phase_events': phase_events,
    'starting_values': start_values
}

with open('seeddata_full.json', 'w', encoding='utf-8') as f:
    json.dump(data_export, f, ensure_ascii=False, indent=2)

print("Export saved to seeddata_full.json")
print()
print("=== SUMMARY ===")
print(f"Starting values: {len(start_values)}")
print(f"Gears: {len(gears)}")
print(f"Action cards: {len(action_cards)}")
print(f"Phase events: {len(phase_events)}")
print()
print("Character mappings:")
for k, v in char_names.items():
    print(f"  {k}: {v} -> {char_refs[k]}")

