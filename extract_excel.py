import openpyxl
import json

wb = openpyxl.load_workbook('Følgeskab_parameter_model.xlsx')

print("Available sheets:", wb.sheetnames)

# Extract all data from sheets
def read_sheet_data(ws, start_row=2):
    headers = [cell.value for cell in ws[1]]
    data = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        if all(v is None for v in row):
            break
        row_dict = {headers[i]: row[i] for i in range(len(headers))}
        data.append(row_dict)
    return data

# Get starting values (6 characters)
start_values = read_sheet_data(wb[wb.sheetnames[1]])
print("=== STARTING VALUES ===")
for person in start_values:
    print(f"{person['ID']}: {person['Person']}")
    print(f"  TjenesteMotivation: {person['Tjenestemotivation']}")
    print(f"  SocialeLyst: {person['Sociale lyst']}")
    print(f"  Tillid: {person['Tillid']}")
    print(f"  Stress: {person['Stress']}")

# Get all Gear impacts (3 phases × 4 gears × 6 people = 72)
gears = read_sheet_data(wb[wb.sheetnames[3]])  # Gear sheet
print("\n=== GEAR (Phase Questions) ===")
gear_by_phase = {}
for gear in gears:
    phase = gear['Fase']
    if phase not in gear_by_phase:
        gear_by_phase[phase] = {}
    gear_num = gear['Gear']
    if gear_num not in gear_by_phase[phase]:
        gear_by_phase[phase][gear_num] = []
    gear_by_phase[phase][gear_num].append(gear)

for phase in ['Opstart', 'Drift', 'Overdragelse']:
    print(f"\n{phase}:")
    if phase in gear_by_phase:
        for gear_num in sorted(gear_by_phase[phase].keys()):
            print(f"  Gear {gear_num}:")
            print(f"    Total impacts: {len(gear_by_phase[phase][gear_num])}")

# Get all Action cards (3 phases × 8 cards × 6 people = 144)
action_cards = read_sheet_data(wb[wb.sheetnames[4]])  # Action cards sheet
print("\n=== ACTION CARDS ===")
cards_by_phase = {}
for card in action_cards:
    phase = card['Fase']
    if phase not in cards_by_phase:
        cards_by_phase[phase] = {}
    card_num = card['Kort']
    if card_num not in cards_by_phase[phase]:
        cards_by_phase[phase][card_num] = []
    cards_by_phase[phase][card_num].append(card)

for phase in ['Opstart', 'Drift', 'Overdragelse']:
    print(f"\n{phase}:")
    if phase in cards_by_phase:
        for card_num in sorted(cards_by_phase[phase].keys()):
            print(f"  Card {card_num}: {len(cards_by_phase[phase][card_num])} people")

total_gear = len(gears)
total_cards = len(action_cards)
print(f"\nTotal Gear impacts: {total_gear}")
print(f"Total Action card impacts: {total_cards}")

# Get phase events
phase_events = read_sheet_data(wb[wb.sheetnames[2]])  # Fasehndelser sheet
print(f"\nTotal Phase events: {len(phase_events)}")

# Save all data to JSON for later use
with open('excel_data.json', 'w', encoding='utf-8') as f:
    json.dump({
        'starting_values': start_values,
        'gears': gears,
        'action_cards': action_cards,
        'phase_events': phase_events
    }, f, ensure_ascii=False, indent=2)

print("\nData saved to excel_data.json")



